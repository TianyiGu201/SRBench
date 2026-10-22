"""Build seed-paired dropout effect sizes and the endpoint-only statistical table.

Input replicate means are produced by build_dropout_replicate_means.py using
the original per-gene aggregation logic. Replicates 1/2/3 correspond to fixed seeds 123/124/125.
"""
from __future__ import annotations

import argparse
from pathlib import Path
import numpy as np
import pandas as pd
from scipy import stats
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
SOURCE = HERE / 'dropout_replicate_means_all_Xenium.csv'
PAIRED_CSV = HERE / 'dropout_seed_paired_effects.csv'
ALL_CSV = HERE / 'dropout_all_effects_summary.csv'
END_CSV = HERE / 'dropout_endpoint_10_vs_50.csv'
WORKBOOK = HERE / 'Gu_Supplementary_Table_6.xlsx'
ENDPOINT_TEXT = HERE / 'Dataset1_endpoint_results.txt'
DATASETS = ['Xe_hsa_breast_c', 'Xe_hsa_colon_h', 'Xe_hsa_kidney_h', 'Xe_hsa_heart_h', 'Xe_hsa_skin_h', 'Xe_hsa_liver_h']
METHODS = ['XFuse', 'TESLA', 'iSTAR', 'scstGCN', 'SuperST_t', 'SuperST_t_g']
METRICS = ['PCC', 'SCC', 'JSD', 'KLD']
RATIOS = [0.1, 0.2, 0.3, 0.4, 0.5]
SEED_MAP = {1: 123, 2: 124, 3: 125}
N = 3
TCRIT = stats.t.ppf(0.975, N - 1)


def summarize(values):
    vals = np.asarray(values, dtype=float)
    assert vals.size == N and np.isfinite(vals).all()
    mean = float(np.mean(vals))
    sd = float(np.std(vals, ddof=1))
    se = sd / np.sqrt(N)
    return mean, sd, se, mean - TCRIT * se, mean + TCRIT * se


def holm(p_values):
    p = np.asarray(p_values, dtype=float)
    if not np.isfinite(p).all():
        raise ValueError('Holm family contains a nonfinite P value')
    order = np.argsort(p, kind='stable')
    ranked = p[order]
    adjusted_ranked = np.minimum(1.0, np.maximum.accumulate((len(p) - np.arange(len(p))) * ranked))
    adjusted = np.empty_like(adjusted_ranked)
    adjusted[order] = adjusted_ranked
    return adjusted


def write_workbook(all_df, end_df):
    wb = Workbook()
    description = wb.active
    description.title = 'Column_Descriptions'
    descriptions = [
        ('Item', 'Definition'),
        ('Source and experimental unit', 'Each replicate is one fixed-seed dropout simulation; rep1 = seed 123, rep2 = seed 124, and rep3 = seed 125. The same seed is matched across dropout ratios within each dataset–method–metric combination.'),
        ('Dataset', 'Dataset number (1–6); Dataset_ID gives the original dataset identifier.'),
        ('Method and Metric', 'Prediction method and performance measure (PCC, SCC, JSD or KLD). All changes remain on the original metric scale.'),
        ('Reference condition', '10% dropout (Dropout_ratio = 0.1). The 10% paired change is exactly zero by definition.'),
        ('Delta (Δ metric)', 'For each seed s and dropout ratio d: M(d,s) − M(0.1,s). Positive values indicate an increase in the named metric, without implying that every metric has the same favorable direction.'),
        ('N_matched_replicates', 'Number of seeds with both the ratio of interest and the 10% reference available. Formal statistics are reported for all three matched seeds; missing evaluations are marked Not evaluated.'),
        ('Mean_delta, SD_delta and SE_delta', 'Arithmetic mean, sample standard deviation (ddof = 1), and SD_delta / sqrt(3) of the three paired changes.'),
        ('CI95_lower and CI95_upper', 'Two-sided 95% Student t interval for the mean paired change: Mean_delta ± t(0.975, df = 2) × SD_delta / sqrt(3). It is based on within-seed differences, not independent-group variances.'),
        ('Mean_10pct, SD_10pct, Mean_50pct and SD_50pct', 'Mean and sample standard deviation across seeds 123, 124 and 125 at the two endpoint ratios.'),
        ('Raw_P', 'Two-sided paired t-test for the additional 50% minus 10% dropout contrast; equivalent to a one-sample t-test of the three paired differences against zero.'),
        ('Holm_adjusted_P', 'Holm step-down adjusted endpoint P value. Each separate Dataset × Metric family comprises the evaluated methods for that dataset and metric; there is no global adjustment across datasets or metrics.'),
        ('Testing scope', 'Formal testing is restricted to the 10% versus 50% endpoint contrast. The 20%, 30% and 40% levels are descriptive paired effect-size and uncertainty summaries, preserving potentially non-monotonic responses; no separate tests are performed for those levels.'),
        ('Status / missing evaluations', 'Available indicates three matched seeds. Not evaluated is the retained table label for unavailable metric summaries; all 15 scstGCN attempts in Dataset 5 produced degenerate all-zero outputs and were non-evaluable. These attempts were performed, but no performance values are imputed.'),
        ('Numeric precision', 'Calculations retain input floating-point precision. Numeric Excel cells are not rounded; display formats show at least ten decimal places or scientific notation for P values.'),
    ]
    for row in descriptions:
        description.append(row)
    for sheet_name, frame in [('All_Dropout_Effects', all_df), ('Endpoint_10_vs_50', end_df)]:
        sheet = wb.create_sheet(sheet_name)
        sheet.append(list(frame.columns))
        for values in frame.itertuples(index=False, name=None):
            sheet.append([None if pd.isna(v) else (v.item() if isinstance(v, np.generic) else v) for v in values])
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for col_idx, name in enumerate(frame.columns, 1):
            col = get_column_letter(col_idx)
            width = min(34, max(12, len(name) + 2))
            sheet.column_dimensions[col].width = width
            if name in {'Raw_P', 'Holm_adjusted_P'}:
                fmt = '0.000000000000E+00'
            elif name in {'Dropout_ratio'}:
                fmt = '0.0'
            elif name in {'N_matched_replicates'}:
                fmt = '0'
            elif name not in {'Dataset','Dataset_ID','Method','Metric','Status'}:
                fmt = '0.000000000000'
            else:
                continue
            for cell in sheet[col][1:]:
                if isinstance(cell.value, (int,float)):
                    cell.number_format = fmt
    description.column_dimensions['A'].width = 36
    description.column_dimensions['B'].width = 115
    for row in description:
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=cell.column == 2)
        description.row_dimensions[row[0].row].height = (max(30, 15 * int(np.ceil(len(str(row[1].value)) / 105)))
                                                       if row[0].row > 1 else 18)
    description.freeze_panes = 'A2'
    for sheet in wb:
        for cell in sheet[1]:
            cell.font = Font(name='Arial', size=10, bold=True)
            cell.alignment = Alignment(vertical='center')
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name='Arial', size=10)
    wb.save(WORKBOOK)


def main():
    global SOURCE, PAIRED_CSV, ALL_CSV, END_CSV, WORKBOOK, ENDPOINT_TEXT
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input', type=Path, required=True, help='Replicate means CSV')
    parser.add_argument('--output-dir', type=Path, required=True)
    args = parser.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    SOURCE = args.input
    PAIRED_CSV = args.output_dir / PAIRED_CSV.name
    ALL_CSV = args.output_dir / ALL_CSV.name
    END_CSV = args.output_dir / END_CSV.name
    WORKBOOK = args.output_dir / WORKBOOK.name
    ENDPOINT_TEXT = args.output_dir / ENDPOINT_TEXT.name
    source = pd.read_csv(SOURCE, float_precision='round_trip')
    keys = ['dataset','method','metric','dropout_ratio','replicate']
    assert not source.duplicated(keys).any(), 'Duplicate replicate result'
    assert not source[keys + ['replicate_gene_mean']].isna().any().any()
    assert set(source.dataset) == set(DATASETS)
    assert set(source.method) == set(METHODS)
    assert set(source.metric) == set(METRICS)
    assert set(source.replicate) == set(SEED_MAP)
    assert set(source.dropout_ratio) == set(RATIOS)
    source['seed'] = source.replicate.map(SEED_MAP)
    assert not source.seed.isna().any()
    paired_rows, all_rows, endpoint_rows = [], [], []
    for dataset_no, dataset in enumerate(DATASETS, 1):
        for method in METHODS:
            for metric in METRICS:
                subset = source[(source.dataset == dataset) & (source.method == method) & (source.metric == metric)]
                common = {'Dataset': f'Dataset {dataset_no}', 'Dataset_ID': dataset, 'Method': method, 'Metric': metric}
                if subset.empty:
                    assert dataset_no == 5 and method == 'scstGCN'
                    for ratio in RATIOS:
                        all_rows.append({**common, 'Dropout_ratio': ratio, 'N_matched_replicates': None,
                                         'Mean_delta': None, 'SD_delta': None, 'SE_delta': None,
                                         'CI95_lower': None, 'CI95_upper': None, 'Status': 'Not evaluated'})
                    endpoint_rows.append({**common, 'N_matched_replicates': None,
                        'Mean_10pct': None, 'SD_10pct': None, 'Mean_50pct': None, 'SD_50pct': None,
                        'Mean_delta_50_minus_10': None, 'SD_delta': None, 'CI95_lower': None,
                        'CI95_upper': None, 'Raw_P': None, 'Holm_adjusted_P': None, 'Status': 'Not evaluated'})
                    continue
                assert len(subset) == 15 and set(subset.dropout_ratio) == set(RATIOS)
                assert all(set(g.replicate) == set(SEED_MAP) for _, g in subset.groupby('dropout_ratio'))
                pivot = subset.pivot(index='replicate', columns='dropout_ratio', values='replicate_gene_mean').loc[[1,2,3], RATIOS]
                assert np.isfinite(pivot.to_numpy()).all()
                base = pivot[0.1].to_numpy()
                for ratio in RATIOS:
                    endpoint = pivot[ratio].to_numpy()
                    delta = endpoint - base if ratio != 0.1 else np.zeros(N, dtype=float)
                    mean, sd, se, low, high = summarize(delta)
                    all_rows.append({**common, 'Dropout_ratio': ratio, 'N_matched_replicates': N,
                        'Mean_delta': mean, 'SD_delta': sd, 'SE_delta': se,
                        'CI95_lower': low, 'CI95_upper': high, 'Status': 'Available'})
                    for rep, value in zip([1,2,3], delta):
                        paired_rows.append({'Dataset': common['Dataset'], 'Dataset_ID': dataset,
                            'Method': method, 'Metric': metric, 'Dropout_ratio': ratio,
                            'Replicate': rep, 'Seed': SEED_MAP[rep], 'Mean_metric': float(endpoint[rep-1]),
                            'Reference_10pct': float(base[rep-1]), 'Delta': float(value)})
                end = pivot[0.5].to_numpy()
                diff = end - base
                mean, sd, se, low, high = summarize(diff)
                p = float(stats.ttest_rel(end, base, alternative='two-sided').pvalue)
                p_check = float(stats.ttest_1samp(diff, 0, alternative='two-sided').pvalue)
                assert np.isclose(p, p_check, rtol=1e-12, atol=1e-15)
                endpoint_rows.append({**common, 'N_matched_replicates': N,
                    'Mean_10pct': float(base.mean()), 'SD_10pct': float(base.std(ddof=1)),
                    'Mean_50pct': float(end.mean()), 'SD_50pct': float(end.std(ddof=1)),
                    'Mean_delta_50_minus_10': mean, 'SD_delta': sd, 'CI95_lower': low,
                    'CI95_upper': high, 'Raw_P': p, 'Holm_adjusted_P': None, 'Status': 'Available'})
    paired = pd.DataFrame(paired_rows)
    all_df = pd.DataFrame(all_rows)
    end_df = pd.DataFrame(endpoint_rows)
    assert len(paired) == 35*4*5*3 and len(all_df) == 36*4*5 and len(end_df) == 36*4
    assert (paired.loc[paired.Dropout_ratio == 0.1, 'Delta'] == 0).all()
    for (dataset, metric), indices in end_df.loc[end_df.Status == 'Available'].groupby(['Dataset','Metric'], sort=False).groups.items():
        raw = end_df.loc[list(indices), 'Raw_P'].to_numpy(float)
        end_df.loc[list(indices), 'Holm_adjusted_P'] = holm(raw)
        assert len(raw) == (5 if dataset == 'Dataset 5' else 6)
    # Write the design ratios as one-decimal labels while retaining 17-digit
    # precision for all measured values and derived effect sizes.
    paired.assign(Dropout_ratio=paired.Dropout_ratio.map(lambda x: f'{x:.1f}')).to_csv(
        PAIRED_CSV, index=False, float_format='%.17g')
    all_df.assign(Dropout_ratio=all_df.Dropout_ratio.map(lambda x: f'{x:.1f}')).to_csv(
        ALL_CSV, index=False, float_format='%.17g')
    end_df.to_csv(END_CSV, index=False, float_format='%.17g')
    write_workbook(all_df, end_df)
    lines = ['Dataset 1 endpoint contrast: 50% minus 10% dropout, three matched seeds (123, 124, 125).',
             'Two-sided paired t-test; Holm correction across six methods within Dataset 1 × PCC.',
             '']
    for method in ['iSTAR','scstGCN','SuperST_t','SuperST_t_g']:
        r = end_df[(end_df.Dataset == 'Dataset 1') & (end_df.Method == method) & (end_df.Metric == 'PCC')].iloc[0]
        lines.extend([method, f"  Mean at 10%: {r.Mean_10pct:.10f}", f"  Mean at 50%: {r.Mean_50pct:.10f}",
            f"  Paired Δ50–10: {r.Mean_delta_50_minus_10:+.10f}",
            f"  95% CI: [{r.CI95_lower:+.10f}, {r.CI95_upper:+.10f}]",
            f"  Raw P: {r.Raw_P:.10g}", f"  Holm-adjusted P: {r.Holm_adjusted_P:.10g}", ''])
    ENDPOINT_TEXT.write_text('\n'.join(lines), encoding='utf-8')
    print('Saved', WORKBOOK)
    print('Saved', PAIRED_CSV, ALL_CSV, END_CSV, ENDPOINT_TEXT)
    print(f'Available endpoint tests: {sum(end_df.Status == "Available")}; Not evaluated rows: {sum(end_df.Status == "Not evaluated")}')
    print('\n'.join(lines))

if __name__ == '__main__':
    main()
