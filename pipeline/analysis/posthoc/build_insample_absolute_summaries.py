"""Build Supplementary Table 2 from replicate-averaged source workbooks."""
from __future__ import annotations
import argparse
import hashlib
import math
import posixpath
import re
import statistics
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

METHODS = ['XFuse','TESLA','iSTAR','scstGCN','SuperST_t','SuperST_t_g']
METRICS = ['PCC','SCC','KLD','JSD','SSIM','PSNR']
NS = {'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
TAG = '{'+NS['m']+'}'

def digest(path):
    with Path(path).open('rb') as handle:
        return hashlib.file_digest(handle,'sha256').hexdigest()

def read_groups(paths):
    groups, names, provenance = {}, {}, []
    for path in paths:
        sha = digest(path)
        with ZipFile(path) as z:
            rels = {x.attrib['Id']:x.attrib['Target'] for x in ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))}
            strings=[]
            if 'xl/sharedStrings.xml' in z.namelist():
                strings=[''.join(t.text or '' for t in x.iter(TAG+'t')) for x in ET.fromstring(z.read('xl/sharedStrings.xml'))]
            def value(cell):
                if cell.attrib.get('t')=='inlineStr':return ''.join(t.text or '' for t in cell.iter(TAG+'t'))
                v=cell.find(TAG+'v')
                if v is None:return None
                if cell.attrib.get('t')=='s':return strings[int(v.text)]
                if cell.attrib.get('t')=='e':raise ValueError(f'Excel error in {path}: {cell.attrib.get("r")}')
                return v.text
            for sheet in ET.fromstring(z.read('xl/workbook.xml')).find('m:sheets',NS):
                match=re.fullmatch(r'(Fig2c|SuppFig5)_D(\d+)_(.+)',sheet.attrib['name'])
                if not match:continue
                dataset,method=int(match[2]),match[3]
                if method not in METHODS:raise ValueError(f'Unknown method: {method}')
                key=(dataset,method)
                if key in groups:raise ValueError(f'Duplicate source group: {key}')
                target=rels[sheet.attrib['{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id']]
                target=target.lstrip('/') if target.startswith('/') else posixpath.normpath('xl/'+target)
                rows,header,meta=[],None,{}
                with z.open(target) as handle:
                    for _,row in ET.iterparse(handle,events=('end',)):
                        if row.tag!=TAG+'row':continue
                        cells={re.sub(r'\d','',c.attrib['r']):value(c) for c in row if c.tag==TAG+'c'}
                        number=int(row.attrib['r'])
                        if number<=6:
                            for col,val in cells.items():meta[f'{col}{number}']=val
                        if 'gene' in cells.values():header={col:val for col,val in cells.items()}
                        elif header and any(v is not None for v in cells.values()):
                            record={field:cells.get(col) for col,field in header.items()}
                            if record.get('gene') is not None:rows.append(record)
                        row.clear()
                if header is None or not set(METRICS).issubset(set(header.values())):
                    raise ValueError(f'Missing metric columns: {sheet.attrib["name"]}')
                if meta.get('B3')!=f'Dataset{dataset}' or meta.get('B5')!=method:
                    raise ValueError(f'Source metadata mismatch: {sheet.attrib["name"]}')
                name=meta['B4']
                if dataset in names and names[dataset]!=name:raise ValueError('Dataset names disagree')
                names[dataset]=name
                genes=[str(r['gene']) for r in rows]
                if not genes or len(set(genes))!=len(genes) or len({g.casefold() for g in genes})!=len(genes):
                    raise ValueError(f'Empty or duplicate gene identifiers: {key}')
                for r in rows:
                    if r.get('dataset')!=name or r.get('method')!=method:raise ValueError(f'Row metadata mismatch: {key}')
                groups[key]=rows
                provenance.append([f'Dataset{dataset}',name,method,Path(path).name,sheet.attrib['name'],len(rows),sha])
    expected={(d,m) for d in range(1,14) for m in METHODS}
    if set(groups)!=expected:raise ValueError(f'Expected 78 source groups; missing={sorted(expected-set(groups))}, extra={sorted(set(groups)-expected)}')
    return groups,names,provenance

def finite_values(rows,field):
    result=[]
    for row in rows:
        val=row.get(field)
        if val is None or val=='':continue
        try:x=float(val)
        except (TypeError,ValueError):continue
        if math.isfinite(x):result.append(x)
    return result

def summarize(groups,names):
    pcc,other,coverage=[],[],[]
    for d in range(1,14):
        labels={}
        for method in ['iSTAR']+[m for m in METHODS if m!='iSTAR']:
            for row in groups[(d,method)]:labels.setdefault(row['gene'].casefold(),row['gene'])
        total=len(labels)
        for method in METHODS:
            rows=groups[(d,method)];v=finite_values(rows,'PCC')
            if not v or any(x < -1 or x > 1 for x in v):raise ValueError(f'Invalid PCC group: {(d,method)}')
            pcc.append([f'Dataset{d}',names[d],method,total,len(v),total-len(v),math.fsum(v)/len(v),statistics.median(v),max(v),math.fsum(x*x for x in v)/len(v),sum(x>.3 for x in v)/len(v),sum(x<.1 for x in v)/len(v)])
            additional=[f'Dataset{d}',names[d],method,total]
            for metric in METRICS[1:]:
                a=finite_values(rows,metric)
                additional.extend([len(a),math.fsum(a)/len(a) if a else None,statistics.median(a) if a else None])
            other.append(additional)
            present={row['gene'].casefold() for row in rows};missing=sorted(labels[g] for g in set(labels)-present)
            reason=None
            if missing:
                reason='Absent from the supplied source table; confirm the reason from the run record rather than inferring it from the gene name.'
            coverage.append([f'Dataset{d}',names[d],method,len(present),total,len(missing),'; '.join(missing) if missing else 'None',reason])
    return pcc,other,coverage

def build(paths,output):
    if Path(output).resolve() in {Path(p).resolve() for p in paths}:
        raise ValueError('Output must not overwrite a source workbook')
    groups,names,provenance=read_groups(paths)
    pcc,other,coverage=summarize(groups,names)
    wb=Workbook();desc=wb.active;desc.title='Column_Descriptions'
    for row in [
        ['Column / item','Description'],
        ['Scope','Post hoc descriptive summaries of the existing registered in-sample predictions (Supplementary Table 2).'],
        ['Replicates','Input gene-level metric values were averaged across three computational runs before summarization across genes. No additional replicate averaging is performed here.'],
        ['N_genes','Case-insensitive union of evaluated gene identifiers across the six method outputs within the dataset. Visium HD summaries describe the evaluated SVG subset.'],
        ['N_valid_PCC / N_missing_PCC','Number of finite gene-level PCC values / N_genes minus that number. Missing source genes are identified separately in Gene_coverage.'],
        ['Mean_squared_PCC','Mean across genes of squared replicate-averaged PCC; not residual-based prediction R² or the square of Mean_PCC.'],
        ['Threshold fractions','Finite PCC > 0.3 or PCC < 0.1 divided by N_valid_PCC. The latter includes negative correlations. These are descriptive thresholds, not universal usability criteria.'],
        ['Other metrics','Unweighted mean and median of finite gene-level values, with separate denominators for each metric. Missing values are not imputed; zero values are retained.'],
        ['Source_provenance','Source workbook, worksheet, gene-row count and workbook SHA-256. Source workbooks are read only.'],
    ]:desc.append(row)
    h1=['Dataset_ID','Dataset_name','Method','N_genes','N_valid_PCC','N_missing_PCC','Mean_PCC','Median_PCC','Maximum_PCC','Mean_squared_PCC','Fraction_PCC_gt_0.3','Fraction_PCC_lt_0.1']
    h2=['Dataset_ID','Dataset_name','Method','N_genes']
    for m in METRICS[1:]:h2.extend(['N_valid_'+m,'Mean_'+m,'Median_'+m])
    h3=['Dataset_ID','Dataset_name','Method','N_genes_present','N_genes_dataset_total','N_genes_absent_from_source','Genes_absent_from_source','Reason_for_missing_genes']
    for name,title,header,rows in [
        ('PCC_summary','Absolute in-sample performance; current Supplementary Table 2.',h1,pcc),
        ('Other_metrics_summary','Complementary absolute in-sample accuracy summaries.',h2,other),
        ('Gene_coverage','Gene coverage in the final source-data worksheets.',h3,coverage),
    ]:
        ws=wb.create_sheet(name);ws.append([title]);ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=len(header));ws.append(header)
        for row in rows:ws.append(row)
    ws=wb.create_sheet('Source_provenance');ws.append(['Dataset_ID','Dataset_name','Method','Source_workbook','Source_worksheet','N_source_genes','Source_workbook_SHA256'])
    for row in provenance:ws.append(row)
    for ws in wb:
        for row in ws:
            for cell in row:cell.font=Font(name='Arial',size=10);cell.alignment=Alignment(vertical='top',wrap_text=True)
        ws.column_dimensions['A'].width=24;ws.column_dimensions['B'].width=36
    output=Path(output);output.parent.mkdir(parents=True,exist_ok=True);wb.save(output)
    return {'output':str(output),'groups':len(groups),'missing_gene_groups':sum(row[5]>0 for row in coverage)}

def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--fig2-source',type=Path,required=True)
    p.add_argument('--suppfig5-source',type=Path,required=True)
    p.add_argument('--output',type=Path,required=True)
    args=p.parse_args()
    print(build([args.fig2_source,args.suppfig5_source],args.output))

if __name__=='__main__':main()
